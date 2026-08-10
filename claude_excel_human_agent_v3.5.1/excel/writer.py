from openpyxl import load_workbook

def ensure_columns(path,sheet,columns):
    wb=load_workbook(path); ws=wb[sheet]
    existing={c.value for c in ws[1] if c.value}
    next_col=ws.max_column+1
    changed=False
    for col in columns:
        if col not in existing:
            ws.cell(1,next_col).value=col
            existing.add(col)
            next_col+=1
            changed=True
    if changed: wb.save(path)
    wb.close()

class ExcelWriter:
    def __init__(self,path,sheet): self.path,self.sheet=path,sheet
    def _cols(self,ws): return {c.value:c.column for c in ws[1]}
    def write(self,row,values):
        wb=load_workbook(self.path); ws=wb[self.sheet]; c=self._cols(ws)
        for k,v in values.items():
            if k in c: ws.cell(row,c[k]).value=v
        wb.save(self.path); wb.close()
