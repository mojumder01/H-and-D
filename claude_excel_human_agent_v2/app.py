import shutil, time
from pathlib import Path
from openpyxl import load_workbook
from config import INPUT_DIR, OUTPUT_DIR, MAX_RETRIES, BETWEEN_ROWS_MS
from excel.reader import ExcelReader
from excel.writer import ExcelWriter, ensure_columns
from browser.chrome_connector import ChromeConnector

TASKS={
    "hd":{
        "label":"Highlights + Description",
        "status_col":"Status","retry_col":"Retry","error_col":"Error",
        "output_map":{"highlights_html":"Highlights HTML","description_html":"Description HTML"},
    },
    "weight":{
        "label":"Weight",
        "status_col":"Weight Status","retry_col":"Weight Retry","error_col":"Weight Error",
        "output_map":{"weight_kg":"Weight (kg)"},
    },
}

def select_excel():
    files=sorted([p for p in Path(INPUT_DIR).glob("*") if p.suffix.lower() in (".xlsx",".xlsm") and not p.name.startswith("~$")])
    if not files:
        print("No Excel file found in input folder.")
        return None
    print("\nAvailable Excel files:")
    for i,f in enumerate(files,1): print(f"  [{i}] {f.name}")
    while True:
        try:
            n=int(input(f"Select file [1-{len(files)}]: "))
            if 1<=n<=len(files): return files[n-1]
        except ValueError: pass
        print("Invalid selection.")

def select_sheet(path):
    wb=load_workbook(path,read_only=True); sheets=wb.sheetnames; wb.close()
    if len(sheets)==1: return sheets[0]
    print("\nAvailable sheets:")
    for i,s in enumerate(sheets,1): print(f"  [{i}] {s}")
    while True:
        try:
            n=int(input(f"Select sheet [1-{len(sheets)}]: "))
            if 1<=n<=len(sheets): return sheets[n-1]
        except ValueError: pass
        print("Invalid selection.")

def select_tasks():
    print("\nWhat do you want to run?")
    print("  [1] Highlights + Description")
    print("  [2] Weight")
    print("  [3] Both (Highlights + Description + Weight)")
    while True:
        try:
            n=int(input("Select [1-3]: "))
            if n==1: return {"hd"}
            if n==2: return {"weight"}
            if n==3: return {"hd","weight"}
        except ValueError: pass
        print("Invalid selection.")

def needed_columns(tasks):
    cols=[]
    for t in tasks:
        cfg=TASKS[t]
        cols+=[cfg["status_col"],cfg["retry_col"],cfg["error_col"]]
        cols+=list(cfg["output_map"].values())
    return cols

def remaining_tasks(row,tasks):
    return {t for t in tasks if str(row.get(TASKS[t]["status_col"]) or "").upper()!="COMPLETED"}

def status_values(tasks,status,attempt,error):
    values={}
    for t in tasks:
        cfg=TASKS[t]
        values[cfg["status_col"]]=status
        values[cfg["retry_col"]]=attempt
        values[cfg["error_col"]]=error
    return values

def result_values(tasks,result,attempt,note):
    values=status_values(tasks,"COMPLETED",attempt,note)
    for t in tasks:
        for key,col in TASKS[t]["output_map"].items():
            values[col]=result[key]
    return values

def main():
    print("="*65)
    print(" Claude Excel Human Automation Agent v2")
    print(" Existing Chrome Session | No API Key")
    print("="*65)
    input_file=select_excel()
    if not input_file: return
    sheet=select_sheet(input_file)
    tasks=select_tasks()

    output=Path(OUTPUT_DIR)/(input_file.stem+"_processed"+input_file.suffix)
    Path(OUTPUT_DIR).mkdir(exist_ok=True)
    shutil.copy2(input_file,output)
    ensure_columns(str(output),sheet,needed_columns(tasks))

    print("\nPaste the Claude conversation URL.")
    print("Example: https://claude.ai/chat/xxxxxxxx-xxxx-xxxx-xxxx-xxxxxxxxxxxx")
    print("Or use a fresh chat: https://claude.ai/new")
    while True:
        url=input("Conversation URL: ").strip()
        if url.startswith("https://claude.ai/"): break
        print("Invalid Claude URL.")

    reader=ExcelReader(str(output),sheet)
    writer=ExcelWriter(str(output),sheet)
    rows=reader.read_rows()
    print(f"\nTotal rows: {len(rows)}")

    connector=ChromeConnector()
    agent=connector.open_conversation(url)

    try:
        for i,row in enumerate(rows,1):
            r=row["_excel_row"]
            remaining=remaining_tasks(row,tasks)
            if not remaining:
                print(f"[{i}/{len(rows)}] SKIP row {r}"); continue
            prior_errors=[]
            for attempt in range(1,MAX_RETRIES+1):
                try:
                    writer.write(r,status_values(remaining,"PROCESSING",attempt,""))
                    result=agent.process(row,remaining)
                    note="; ".join(f"attempt {n} failed: {msg}" for n,msg in prior_errors)
                    writer.write(r,result_values(remaining,result,attempt,note))
                    print(f"[{i}/{len(rows)}] DONE row {r}")
                    break
                except Exception as e:
                    msg=str(e)
                    print(f"[{i}/{len(rows)}] Attempt {attempt} failed: {msg}")
                    prior_errors.append((attempt,msg))
                    writer.write(r,status_values(remaining,"RETRYING",attempt,msg))
                    if attempt==MAX_RETRIES:
                        writer.write(r,status_values(remaining,"FAILED",attempt,msg))
                    time.sleep(2**attempt)
            time.sleep(BETWEEN_ROWS_MS/1000)
    finally:
        connector.close()

    print("\nFinished:",output)

if __name__=="__main__": main()
