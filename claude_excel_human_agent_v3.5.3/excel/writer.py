import os
from openpyxl import load_workbook

def wait_until_unlocked(fn,file_label):
    # On Windows, opening the output file in Excel (or antivirus/OneDrive
    # briefly scanning it) blocks other processes from overwriting/renaming
    # it, raising PermissionError. For a run that can take multiple days,
    # someone opening the file to peek at progress is a near-certainty -
    # that must not crash the whole job. Wait for the user to close it and
    # retry, rather than losing the row/state that was about to be saved.
    while True:
        try:
            return fn()
        except PermissionError:
            print(f"\nCould not save - '{file_label}' appears to be open in Excel "
                  f"(or locked by another program). Close it, then press Enter to retry...")
            input()

def _atomic_save(wb,path):
    # If power is cut / the process is killed mid-save, the ORIGINAL file
    # must survive intact - this project has to be resumable across
    # multi-day runs. Saving straight to `path` risks a half-written,
    # corrupted .xlsx if interrupted mid-write. Saving to a temp file first
    # and only then swapping it in with an atomic rename means the on-disk
    # file is always either the old complete version or the new complete
    # version, never a partial one.
    tmp=path+".tmp"
    wb.save(tmp)
    wb.close()
    wait_until_unlocked(lambda: os.replace(tmp,path),os.path.basename(path))

def ensure_columns(path,sheet,columns):
    wb=load_workbook(path); ws=wb[sheet]
    existing={c.value for c in ws[1] if c.value}
    next_col=ws.max_column+1
    changed=False
    for col in columns:
        if col not in existing:
            ws.cell(1,next_col).value=col
            existing.add(col)
            next_col+=1
            changed=True
    if changed: _atomic_save(wb,path)
    else: wb.close()

class ExcelWriter:
    def __init__(self,path,sheet): self.path,self.sheet=path,sheet
    def _cols(self,ws): return {c.value:c.column for c in ws[1]}
    def write(self,row,values):
        wb=load_workbook(self.path); ws=wb[self.sheet]; c=self._cols(ws)
        for k,v in values.items():
            if k in c: ws.cell(row,c[k]).value=v
        _atomic_save(wb,self.path)
