import shutil, time
from pathlib import Path
from openpyxl import load_workbook
from config import INPUT_DIR, OUTPUT_DIR, MAX_RETRIES, BETWEEN_ROWS_MS
from excel.reader import ExcelReader
from excel.writer import ExcelWriter
from browser.chrome_connector import ChromeConnector

def select_excel():
    files=sorted([p for p in Path(INPUT_DIR).glob("*") if p.suffix.lower() in (".xlsx",".xlsm") and not p.name.startswith("~$")])
    if not files:
        print("No Excel file found in input folder.")
        return None
    print("\nAvailable Excel files:")
    for i,f in enumerate(files,1): print(f"  [{i}] {f.name}")
    while True:
        try:
            n=int(input(f"Select file [1-{len(files)}]: "))
            if 1<=n<=len(files): return files[n-1]
        except ValueError: pass
        print("Invalid selection.")

def select_sheet(path):
    wb=load_workbook(path,read_only=True); sheets=wb.sheetnames; wb.close()
    if len(sheets)==1: return sheets[0]
    print("\nAvailable sheets:")
    for i,s in enumerate(sheets,1): print(f"  [{i}] {s}")
    while True:
        try:
            n=int(input(f"Select sheet [1-{len(sheets)}]: "))
            if 1<=n<=len(sheets): return sheets[n-1]
        except ValueError: pass
        print("Invalid selection.")

def main():
    print("="*65)
    print(" Claude Excel Human Automation Agent v1")
    print(" Existing Chrome Session | No API Key")
    print("="*65)
    input_file=select_excel()
    if not input_file: return
    sheet=select_sheet(input_file)

    output=Path(OUTPUT_DIR)/(input_file.stem+"_processed"+input_file.suffix)
    Path(OUTPUT_DIR).mkdir(exist_ok=True)
    shutil.copy2(input_file,output)

    print("\nPaste the Claude conversation URL.")
    print("Example: https://claude.ai/chat/xxxxxxxx-xxxx-xxxx-xxxx-xxxxxxxxxxxx")
    print("Or use a fresh chat: https://claude.ai/new")
    while True:
        url=input("Conversation URL: ").strip()
        if url.startswith("https://claude.ai/"): break
        print("Invalid Claude URL.")

    reader=ExcelReader(str(output),sheet)
    writer=ExcelWriter(str(output),sheet)
    rows=reader.read_rows()
    print(f"\nTotal rows: {len(rows)}")

    connector=ChromeConnector()
    agent=connector.open_conversation(url)

    try:
        for i,row in enumerate(rows,1):
            r=row["_excel_row"]
            if str(row.get("Status") or "").upper()=="COMPLETED":
                print(f"[{i}/{len(rows)}] SKIP row {r}"); continue
            for attempt in range(1,MAX_RETRIES+1):
                try:
                    writer.write_status(r,"PROCESSING",attempt,"")
                    result=agent.process(row)
                    writer.write_result(r,result["highlights_html"],result["description_html"],"COMPLETED",attempt,"")
                    print(f"[{i}/{len(rows)}] DONE row {r}")
                    break
                except Exception as e:
                    print(f"[{i}/{len(rows)}] Attempt {attempt} failed: {e}")
                    writer.write_status(r,"RETRYING",attempt,str(e))
                    if attempt==MAX_RETRIES:
                        writer.write_status(r,"FAILED",attempt,str(e))
                    time.sleep(2**attempt)
            time.sleep(BETWEEN_ROWS_MS/1000)
    finally:
        connector.close()

    print("\nFinished:",output)

if __name__=="__main__": main()
