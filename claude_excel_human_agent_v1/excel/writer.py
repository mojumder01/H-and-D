from openpyxl import load_workbook
class ExcelWriter:
    def __init__(self,path,sheet): self.path,self.sheet=path,sheet
    def _cols(self,ws): return {c.value:c.column for c in ws[1]}
    def write_status(self,row,status,retry,error):
        wb=load_workbook(self.path); ws=wb[self.sheet]; c=self._cols(ws)
        for k,v in {"Status":status,"Retry":retry,"Error":error}.items():
            if k in c: ws.cell(row,c[k]).value=v
        wb.save(self.path); wb.close()
    def write_result(self,row,h,d,status,retry,error):
        wb=load_workbook(self.path); ws=wb[self.sheet]; c=self._cols(ws)
        for k,v in {"Highlights HTML":h,"Description HTML":d,"Status":status,"Retry":retry,"Error":error}.items():
            if k in c: ws.cell(row,c[k]).value=v
        wb.save(self.path); wb.close()
