from openpyxl import load_workbook
class ExcelReader:
    def __init__(self,path,sheet): self.path,self.sheet=path,sheet
    def read_rows(self):
        wb=load_workbook(self.path); ws=wb[self.sheet]
        headers=[c.value for c in ws[1]]; rows=[]
        for n,row in enumerate(ws.iter_rows(min_row=2),2):
            d={h:c.value for h,c in zip(headers,row)}; d["_excel_row"]=n; rows.append(d)
        wb.close(); return rows
